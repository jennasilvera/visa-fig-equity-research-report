from pathlib import Path
from openpyxl import load_workbook

MODEL_PATH = Path("model/Visa_Institutional_Model_v2.xlsx")
REPORT_PATH = Path("model/model_checks/model_validation_report.md")

required_sheets = [
    "01_Cover",
    "02_Sources",
    "03_Setup_Control",
    "04_Historical_Financials",
    "05_KPI_History",
    "06_Revenue_Build",
    "07_Client_Incentives",
    "08_VAS_Model",
    "09_Cross_Border_Model",
    "10_Operating_Expense_Build",
    "11_Margin_Walk",
    "12_Tax_DA_Capex_NWC",
    "13_Free_Cash_Flow",
    "14_Share_Count_Buybacks",
    "15_DCF",
    "16_Trading_Comps",
    "17_Precedent_Transactions",
    "18_SOTP",
    "19_Scenario_Manager",
    "20_Sensitivity_Tables",
    "21_Consensus_Bridge",
    "22_Output_Dashboard",
    "23_Charts",
    "24_Model_Checks",
    "25_Print_Package",
]

required_files = [
    "data/processed/visa_historical_financials_source_backed.csv",
    "data/processed/visa_operating_kpis_source_backed.csv",
    "data/processed/visa_historical_data_sources.md",
    "model/model_outputs/model_population_summary.csv",
    "scripts/populate_institutional_model_v2.py",
]

def yes_no(condition):
    return "FOUND" if condition else "MISSING"

def main():
    wb = load_workbook(MODEL_PATH, data_only=False)
    sheet_names = wb.sheetnames

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)

    with REPORT_PATH.open("w", encoding="utf-8") as f:
        f.write("# Visa Institutional Model Validation Report\n\n")
        f.write("## Purpose\n\n")
        f.write(
            "This report validates that the institutional Visa model contains the expected workbook tabs, "
            "source-backed data files, and model population outputs.\n\n"
        )

        f.write("## Workbook Status\n\n")
        f.write(f"- Model file: `{MODEL_PATH}`\n")
        f.write(f"- Workbook opens successfully: `YES`\n")
        f.write(f"- Sheet count: `{len(sheet_names)}`\n\n")

        f.write("## Required Sheet Check\n\n")
        f.write("| Sheet | Status |\n")
        f.write("|---|---|\n")
        for sheet in required_sheets:
            f.write(f"| `{sheet}` | {yes_no(sheet in sheet_names)} |\n")

        f.write("\n## Required Source / Output File Check\n\n")
        f.write("| File | Status |\n")
        f.write("|---|---|\n")
        for file in required_files:
            f.write(f"| `{file}` | {yes_no(Path(file).exists())} |\n")

        f.write("\n## Historical Data Spot Check\n\n")
        ws = wb["04_Historical_Financials"]
        f.write("| Metric | FY2021 | FY2022 | FY2023 | FY2024 | FY2025 |\n")
        f.write("|---|---:|---:|---:|---:|---:|\n")
        f.write(
            f"| Net Revenue | {ws['B5'].value} | {ws['C5'].value} | {ws['D5'].value} | {ws['E5'].value} | {ws['F5'].value} |\n"
        )
        f.write(
            f"| Operating Income | {ws['B8'].value} | {ws['C8'].value} | {ws['D8'].value} | {ws['E8'].value} | {ws['F8'].value} |\n"
        )
        f.write(
            f"| Net Income | {ws['B10'].value} | {ws['C10'].value} | {ws['D10'].value} | {ws['E10'].value} | {ws['F10'].value} |\n"
        )
        f.write(
            f"| Diluted EPS | {ws['B11'].value} | {ws['C11'].value} | {ws['D11'].value} | {ws['E11'].value} | {ws['F11'].value} |\n"
        )

        f.write("\n## Model Notes\n\n")
        f.write(
            "- The V2 model has been populated with source-backed FY2021–FY2025 historical financials and operating KPIs.\n"
        )
        f.write(
            "- Historical operating expenses are currently plugged into G&A / Other until a detailed expense breakout is added.\n"
        )
        f.write(
            "- Forecast years, DCF outputs, trading comps, SOTP, and scenario outputs should be refreshed with current market data before publication or interview use.\n"
        )

    print(f"Wrote {REPORT_PATH}")

if __name__ == "__main__":
    main()
