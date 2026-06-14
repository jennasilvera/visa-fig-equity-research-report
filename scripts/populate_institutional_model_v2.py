from pathlib import Path
import csv
from openpyxl import load_workbook

MODEL_PATH = Path("model/Visa_Institutional_Model_v2.xlsx")
FIN_PATH = Path("data/processed/visa_historical_financials_source_backed.csv")
KPI_PATH = Path("data/processed/visa_operating_kpis_source_backed.csv")
OUTPUT_SUMMARY = Path("model/model_outputs/model_population_summary.csv")

YEARS = ["FY2021", "FY2022", "FY2023", "FY2024", "FY2025"]
COLS = {year: idx + 2 for idx, year in enumerate(YEARS)}  # B:F


def read_csv_by_year(path):
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return {row["Fiscal Year"].strip(): row for row in reader}


def parse_value(value):
    if value is None:
        return None
    value = str(value).strip()
    if value == "":
        return None
    if value.endswith("%"):
        return float(value[:-1]) / 100
    return float(value.replace(",", ""))


def set_percent(cell):
    cell.number_format = "0.0%"


def set_number(cell):
    cell.number_format = "#,##0.0"


def set_integer(cell):
    cell.number_format = "#,##0"


def main():
    if not MODEL_PATH.exists():
        raise FileNotFoundError(f"Missing model: {MODEL_PATH}")

    fin = read_csv_by_year(FIN_PATH)
    kpis = read_csv_by_year(KPI_PATH)

    wb = load_workbook(MODEL_PATH)

    # 04 Historical Financials
    ws = wb["04_Historical_Financials"]

    hist_rows = {
        "Net Revenue": 5,
        "Operating Expenses": 7,
        "Operating Income": 8,
        "Operating Margin": 9,
        "Net Income": 10,
        "Diluted EPS": 11,
        "Operating Cash Flow": 12,
        "Capital Expenditures": 13,
        "Free Cash Flow": 14,
    }

    for year, col in COLS.items():
        row = fin[year]
        for field, excel_row in hist_rows.items():
            ws.cell(excel_row, col, parse_value(row[field]))

        # Calculate diluted shares from Net Income / Diluted EPS
        net_income = parse_value(row["Net Income"])
        eps = parse_value(row["Diluted EPS"])
        ws.cell(15, col, net_income / eps if eps else None)

        if col > 2:
            ws.cell(6, col, f"={ws.cell(5, col).coordinate}/{ws.cell(5, col-1).coordinate}-1")

        set_percent(ws.cell(6, col))
        set_percent(ws.cell(9, col))
        set_integer(ws.cell(5, col))
        set_integer(ws.cell(7, col))
        set_integer(ws.cell(8, col))
        set_integer(ws.cell(10, col))
        set_number(ws.cell(11, col))
        set_integer(ws.cell(12, col))
        set_integer(ws.cell(13, col))
        set_integer(ws.cell(14, col))
        set_number(ws.cell(15, col))

    # 06 Revenue Build
    ws = wb["06_Revenue_Build"]

    revenue_rows = {
        "Service Revenue": 5,
        "Data Processing Revenue": 6,
        "International Transaction Revenue": 7,
        "Other Revenue": 8,
    }

    for year, col in COLS.items():
        row = fin[year]
        for field, excel_row in revenue_rows.items():
            ws.cell(excel_row, col, parse_value(row[field]))

        # Client incentives are negative in source CSV, but this model subtracts the value,
        # so store client incentives as a positive contra-revenue number.
        ws.cell(10, col, abs(parse_value(row["Client Incentives"])))

        ws.cell(9, col, f"=SUM({ws.cell(5, col).coordinate}:{ws.cell(8, col).coordinate})")
        ws.cell(11, col, f"={ws.cell(9, col).coordinate}-{ws.cell(10, col).coordinate}")

        if col > 2:
            ws.cell(12, col, f"={ws.cell(11, col).coordinate}/{ws.cell(11, col-1).coordinate}-1")

        for r in [5, 6, 7, 8, 9, 10, 11]:
            set_integer(ws.cell(r, col))
        set_percent(ws.cell(12, col))

    # 05 KPI History
    ws = wb["05_KPI_History"]

    for year, col in COLS.items():
        kpi = kpis[year]
        fin_row = fin[year]

        payments_volume_t = parse_value(kpi["Payments Volume ($T)"])
        processed_tx_b = parse_value(kpi["Processed Transactions (B)"])
        credentials_b = parse_value(kpi["Payment Credentials (B)"])

        net_revenue = parse_value(fin_row["Net Revenue"])
        service = parse_value(fin_row["Service Revenue"])
        data_proc = parse_value(fin_row["Data Processing Revenue"])
        intl = parse_value(fin_row["International Transaction Revenue"])
        other = parse_value(fin_row["Other Revenue"])
        gross_revenue = service + data_proc + intl + other
        incentives = abs(parse_value(fin_row["Client Incentives"]))
        vas_revenue = parse_value(fin_row["VAS Revenue"])

        ws.cell(5, col, payments_volume_t)
        ws.cell(6, col, processed_tx_b)
        ws.cell(8, col, credentials_b)
        ws.cell(10, col, gross_revenue / (payments_volume_t * 1_000_000) if payments_volume_t else None)
        ws.cell(11, col, net_revenue / (payments_volume_t * 1_000_000) if payments_volume_t else None)
        ws.cell(12, col, incentives / gross_revenue if gross_revenue else None)
        ws.cell(13, col, vas_revenue)
        ws.cell(14, col, vas_revenue / net_revenue if vas_revenue and net_revenue else None)

        for r in [10, 11, 12, 14]:
            set_percent(ws.cell(r, col))
        for r in [5, 6, 8]:
            set_number(ws.cell(r, col))
        set_integer(ws.cell(13, col))

    # 08 VAS Model
    ws = wb["08_VAS_Model"]

    for year, col in COLS.items():
        row = fin[year]
        vas_revenue = parse_value(row["VAS Revenue"])
        net_revenue = parse_value(row["Net Revenue"])

        ws.cell(5, col, vas_revenue)
        ws.cell(7, col, f"='06_Revenue_Build'!{ws.cell(11, col).coordinate}")
        ws.cell(8, col, f"={ws.cell(5, col).coordinate}/{ws.cell(7, col).coordinate}")

        if col > 2:
            ws.cell(6, col, f"={ws.cell(5, col).coordinate}/{ws.cell(5, col-1).coordinate}-1")

        set_integer(ws.cell(5, col))
        set_percent(ws.cell(6, col))
        set_integer(ws.cell(7, col))
        set_percent(ws.cell(8, col))

    # 09 Cross-Border Model
    ws = wb["09_Cross_Border_Model"]

    for year, col in COLS.items():
        row = fin[year]
        ws.cell(7, col, parse_value(row["International Transaction Revenue"]))
        if col > 2:
            ws.cell(8, col, f"={ws.cell(7, col).coordinate}/{ws.cell(7, col-1).coordinate}-1")
        set_integer(ws.cell(7, col))
        set_percent(ws.cell(8, col))

    # 10 Operating Expense Build
    ws = wb["10_Operating_Expense_Build"]

    for year, col in COLS.items():
        row = fin[year]

        # Historical opex is placed in G&A / Other as a simplified historical plug.
        # Forward model should break this into actual expense categories later.
        ws.cell(10, col, parse_value(row["Operating Expenses"]))
        ws.cell(11, col, f"=SUM({ws.cell(6, col).coordinate}:{ws.cell(10, col).coordinate})")
        ws.cell(12, col, f"={ws.cell(5, col).coordinate}-{ws.cell(11, col).coordinate}")
        ws.cell(13, col, f"={ws.cell(12, col).coordinate}/{ws.cell(5, col).coordinate}")

        set_integer(ws.cell(10, col))
        set_integer(ws.cell(11, col))
        set_integer(ws.cell(12, col))
        set_percent(ws.cell(13, col))

    # 07 Client Incentives already links to 06_Revenue_Build.
    # Add source note.
    ws = wb["02_Sources"]
    next_row = ws.max_row + 2
    source_rows = [
        ["Company Filing", "Visa annual reports / Form 10-K", "FY2021-FY2025", "Historical financials", "data/processed/visa_historical_financials_source_backed.csv", "Populated into model by script"],
        ["Company Filing", "Visa annual reports / Form 10-K", "FY2021-FY2025", "Operating KPIs", "data/processed/visa_operating_kpis_source_backed.csv", "Populated into model by script"],
    ]
    for r_idx, values in enumerate(source_rows, start=next_row):
        for c_idx, value in enumerate(values, start=1):
            ws.cell(r_idx, c_idx, value)

    wb.save(MODEL_PATH)

    OUTPUT_SUMMARY.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_SUMMARY.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Item", "Status"])
        writer.writerow(["Historical financials populated", "Yes"])
        writer.writerow(["Operating KPIs populated", "Yes"])
        writer.writerow(["Revenue build populated", "Yes"])
        writer.writerow(["Client incentives linked", "Yes"])
        writer.writerow(["VAS model partially populated", "Yes"])
        writer.writerow(["Cross-border revenue populated", "Yes"])
        writer.writerow(["Note", "Historical operating expenses are plugged into G&A / Other until detailed expense breakout is added"])

    print(f"Updated {MODEL_PATH}")
    print(f"Wrote {OUTPUT_SUMMARY}")


if __name__ == "__main__":
    main()
